"""Export the GLO note corpus + analysis into a multi-sheet .xlsx workbook."""
import json, collections, re, statistics
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

REPO = "/Users/jehualli/substack-intelligence"
OUT = f"{REPO}/glo/glo_notes.xlsx"
U = 273636371

notes = json.load(open(f"{REPO}/substack_notes_all.json"))
notes.sort(key=lambda n: n["date"], reverse=True)
orig = [n for n in notes if n["author_id"] == U]

ACCENT = "FF6719"
HEAD_FILL = PatternFill("solid", fgColor="1A1A1A")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1A1A1A")
thin = Side(style="thin", color="E7E7E7")
BORDER = Border(bottom=thin)

wb = Workbook()


def style_header(ws, row=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[row].height = 22
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# --- Sheet 1: all notes ----------------------------------------------------
ws = wb.active
ws.title = "Notes"
cols = ["Date", "Time (UTC)", "Type", "Author", "Handle",
        "Reactions", "Restacks", "Replies", "Chars", "Body", "Note ID"]
ws.append(cols)
for n in notes:
    d = n["date"]
    body = n["body"] or ""
    ws.append([
        d[:10], d[11:19],
        "Original" if n["author_id"] == U else "Restack",
        "GLO" if n["author_id"] == U else (n.get("author_name") or ""),
        "thepillarsofwisdom" if n["author_id"] == U else (n.get("author_handle") or ""),
        n["reactions"], n["restacks"], n["replies"], len(body),
        # Excel hard-caps a cell at 32,767 chars; nothing here is close, but guard anyway.
        body[:32000] if body else "[image-only note]",
        n["id"],
    ])
style_header(ws)
widths(ws, {"A": 12, "B": 11, "C": 10, "D": 22, "E": 20, "F": 11,
            "G": 10, "H": 9, "I": 8, "J": 100, "K": 12})
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=10).alignment = Alignment(wrap_text=False, vertical="top")
ws.auto_filter.ref = f"A1:K{ws.max_row}"

# --- Sheet 2: monthly summary ----------------------------------------------
ws = wb.create_sheet("Monthly")
ws.append(["Month", "Original notes", "Restacks", "Total reactions",
           "Mean reactions", "Median", "Max", "Top note"])
months = sorted({n["date"][:7] for n in notes}, reverse=True)
for m in months:
    o = [n for n in orig if n["date"][:7] == m]
    r = [n for n in notes if n["date"][:7] == m and n["author_id"] != U]
    if not o:
        continue
    rx = [n["reactions"] for n in o]
    top = max(o, key=lambda n: n["reactions"])
    ws.append([m, len(o), len(r), sum(rx), round(statistics.mean(rx), 1),
               statistics.median(rx), max(rx), (top["body"] or "[image]")[:180]])
style_header(ws)
widths(ws, {"A": 10, "B": 14, "C": 10, "D": 15, "E": 14, "F": 9, "G": 8, "H": 90})

# --- Sheet 3: top notes ----------------------------------------------------
ws = wb.create_sheet("Top notes")
ws.append(["Rank", "Date", "Reactions", "Restacks", "Replies", "Body"])
for i, n in enumerate(sorted(orig, key=lambda x: -x["reactions"])[:200], 1):
    ws.append([i, n["date"][:10], n["reactions"], n["restacks"], n["replies"],
               (n["body"] or "[image-only note]")[:2000]])
style_header(ws)
widths(ws, {"A": 7, "B": 12, "C": 11, "D": 10, "E": 9, "F": 110})

# --- Sheet 4: opponent cycles ----------------------------------------------
TARGETS = {
    "Mainstream right / grifters": r"jordan peterson|ben shapiro|charlie kirk|bet[- ]david|nick fuentes|andrew tate|jack murphy|matt walsh|tim pool|cuckservative|conservatard|cuckcervative|controlled opposition|mainstream conservative|conservaslop|rightoid|woke right|daily wire|prager|crowder",
    "Liberal women / feminism": r"liberal women|white wom|feminist|foid|\bfoids\b|shitlib wom|karen|girlboss|onlyfans|\bslut|hypergam|mommies|that girl|metoo|thot",
    "Boomers": r"boomer",
    "Jews / Zionism": r"\bjew|zionis|semit|israel|isreal|talmud|rothschild|masonic|freemason|khazar",
    "Immigration / H-1B / Indians": r"h1b|h-1b|immigrant|migrant|indian|pajeet|jeet|visa|open border|great replacement|third world",
    "Globalists / capital": r"globalist|billionaire|ruling class|zio capital|blackrock|wef|davos|federal reserve|agenda 2030|bankers|financializ|capital needs",
    "Normies / bugmen / NPCs": r"\bnormie|normtard|bugman|bugmen|\bnpc\b|goyim|goycattle|midwit|w2 employ|cubicle|wagie",
    "Leftists / liberals (generic)": r"leftoid|libtard|\bliberals\b|\bliberal \b|shitlib|marxist|communist|progressive|woke\b|dei\b|critical race",
    "State / surveillance / institutions": r"\bfbi\b|\bcia\b|\bfed\b|feds\b|surveillance|chat control|flock camera|digital id|police|government|public education|university|the eu\b|european union|censor",
    "Food industry / GoySlop": r"goyslop|goy slop|seed oil|glycosulphate|glyphosate|canola|processed food|slopified|normie slop|goysnack",
    "Big Tech / AI": r"\bai\b|openai|anthropic|chatgpt|sam altman|\bzuck|silicon valley|algorithm|vibe cod",
}
C = {k: re.compile(v, re.I) for k, v in TARGETS.items()}
asc = sorted({n["date"][:7] for n in orig})
mn = collections.Counter(n["date"][:7] for n in orig)

ws = wb.create_sheet("Opponent cycles")
ws["A1"] = "Share of each month's original notes matching each target category (%)"
ws["A1"].font = TITLE_FONT
ws["A2"] = ("Categories overlap — a note can match several — so columns exceed 100%. "
            "Percentages, not counts: monthly volume ranges 73–264.")
ws["A2"].font = Font(italic=True, size=10, color="757575")

ws.append([])
ws.append(["Target"] + asc)
for k, rx in C.items():
    ws.append([k] + [round(sum(1 for n in orig if n["date"][:7] == m and rx.search(n["body"])) / mn[m] * 100)
                     for m in asc])
ws.append(["Notes that month"] + [mn[m] for m in asc])
style_header(ws, row=4)
ws.freeze_panes = "B5"
widths(ws, {"A": 34})
for i in range(2, len(asc) + 2):
    ws.column_dimensions[get_column_letter(i)].width = 7
# Highlight each category's own peak month.
peak = PatternFill("solid", fgColor="FFE3D2")
for r in range(5, 5 + len(C)):
    vals = [(ws.cell(row=r, column=c).value, c) for c in range(2, len(asc) + 2)]
    mx = max(v for v, _ in vals)
    for v, c in vals:
        if v == mx:
            ws.cell(row=r, column=c).fill = peak
            ws.cell(row=r, column=c).font = Font(bold=True)
ws.cell(row=4 + len(C) + 1, column=1).font = Font(bold=True)

# --- Sheet 5: readme -------------------------------------------------------
ws = wb.create_sheet("About", 0)
ws["A1"] = "GLO — Substack notes corpus"
ws["A1"].font = Font(bold=True, size=16)
rows = [
    ("Source", "substack-intelligence/substack_notes_all.json"),
    ("Profile", "GLO · @thepillarsofwisdom · user 273636371"),
    ("Exported", "2026-07-28"),
    ("Range", f"{notes[-1]['date'][:10]} → {notes[0]['date'][:10]}"),
    ("Total notes", len(notes)),
    ("Original", len(orig)),
    ("Restacks", len(notes) - len(orig)),
    ("", ""),
    ("Notes", "Every note, newest first. Filterable; 'Type' separates original from restacked."),
    ("Monthly", "Volume and engagement per month."),
    ("Top notes", "The 200 highest-reaction original notes."),
    ("Opponent cycles", "Regex target categories by month — see glo_opponent_cycles.md for method and caveats."),
    ("", ""),
    ("Caveat", "Reaction counts are as of the scrape date and keep moving; older notes have had longer to accumulate."),
    ("Caveat", "Body text is reproduced verbatim, including the author's typos and slurs."),
]
for i, (k, v) in enumerate(rows, start=3):
    ws.cell(row=i, column=1, value=k).font = Font(bold=True, size=11)
    ws.cell(row=i, column=2, value=v)
widths(ws, {"A": 18, "B": 96})

wb.save(OUT)
print(f"wrote {OUT}")
print(f"sheets: {wb.sheetnames}")
